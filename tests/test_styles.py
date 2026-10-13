"""Style inheritance, and the false positive that motivated it.

The regression this file guards is concrete: checking ``w:outlineLvl`` on the
paragraph alone reported 40 bogus warnings on a healthy 190-paragraph thesis,
because Word cascades the value from ``styles.xml`` through ``basedOn``. If
inheritance resolution ever regresses, these tests fail loudly.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from docqc.document import Document
from docqc.styles import StyleTable


def make_heading_doc(path: Path, *, levels: int = 3) -> Path:
    """A document whose headings rely on style-level outline levels.

    ``add_heading`` applies Heading N *without* writing a local ``w:outlineLvl``,
    which is exactly the shape that triggered the false positive.

    Args:
        path: Where to write the file.
        levels: How many heading levels to create.

    Returns:
        The written path.
    """
    doc = DocxDocument()
    doc.add_paragraph("正文开头。")
    for level in range(1, levels + 1):
        doc.add_heading(f"第 {level} 级标题", level=level)
        doc.add_paragraph(f"第 {level} 级下面的正文。")
    doc.save(path)
    return path


# -- style table --------------------------------------------------------------


def test_style_table_reads_outline_levels(tmp_path: Path) -> None:
    """Enough of styles.xml is parsed to answer heading questions.

    The default python-docx template always defines Heading1..Heading9, so a
    fresh document is a sufficient source of truth here.
    """
    path = make_heading_doc(tmp_path / "doc.docx")
    table = StyleTable.from_docx(path)
    assert len(table) > 0, "expected the default template's styles to be parsed"
    assert "Heading1" in table
    assert table.get("Heading1") is not None
    assert table.get("Heading1").outline_level == 0  # type: ignore[union-attr]


def test_style_table_tolerates_a_missing_part(tmp_path: Path) -> None:
    """A document with no styles part yields an empty table, not an error."""
    path = tmp_path / "bare.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    table = StyleTable.from_docx(path)
    assert len(table) == 0
    assert table.get("Heading1") is None
    assert table.resolve("Heading1", "outline_level") is None


def test_resolve_rejects_unknown_property(tmp_path: Path) -> None:
    """Resolving a property that is not inheritable is a programming error."""
    table = StyleTable({})
    with pytest.raises(ValueError, match="not inheritable"):
        table.resolve("Heading1", "font_size")


# -- the regression -----------------------------------------------------------


def test_heading_outline_level_is_inherited_not_local(tmp_path: Path) -> None:
    """The paragraph carries no local level, yet resolution must find one.

    This asserts the exact shape that caused the false positive: if this ever
    returns ``None``, the structure checker will start firing on healthy files
    again.
    """
    path = make_heading_doc(tmp_path / "doc.docx")
    document = Document(path)
    headings = [p for p in document.paragraphs if p.style_name_id == "Heading1"]
    assert headings, "fixture must contain at least one Heading1 paragraph"

    para = headings[0]
    assert para.outline_level is None, (
        "python-docx should not write a local outline level; "
        "if it starts doing so this test no longer proves inheritance"
    )

    resolved = document.resolve_style(para)
    assert resolved.outline_level == 0
    assert resolved.source is not None and resolved.source.startswith("style:")


def test_healthy_headings_produce_no_warnings(tmp_path: Path) -> None:
    """The regression, end to end: no false warnings on a normal thesis shape."""
    from docqc.runner import check

    path = make_heading_doc(tmp_path / "doc.docx", levels=3)
    report = check(path, names=["structure"])
    offenders = [
        f.code
        for f in report.findings()
        if f.severity.value in {"error", "warning"}
        and f.code in {"HEADING_WITHOUT_OUTLINE_LEVEL", "NO_HEADINGS"}
    ]
    assert offenders == [], f"false positives returned: {offenders}"


def test_heading_level_jump_is_reported(tmp_path: Path) -> None:
    """Skipping a level (1 -> 3) means the outline has a hole."""
    from docqc.runner import check

    doc = DocxDocument()
    doc.add_heading("一级", level=1)
    doc.add_paragraph("正文。")
    doc.add_heading("三级", level=3)
    doc.add_paragraph("正文。")
    path = tmp_path / "jump.docx"
    doc.save(path)

    report = check(path, names=["structure"])
    assert "HEADING_LEVEL_JUMP" in {f.code for f in report.findings()}
