"""Checker behaviour, asserted against generated defect fixtures.

Each test plants exactly one defect and asserts the specific code fires — not
merely that *something* was reported. A checker that reports the wrong code is
as broken as one that reports nothing, and only per-code assertions catch that.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from docqc.checkers import formula as formula_module  # noqa: F401 - registers
from docqc.checkers import layout as layout_module  # noqa: F401 - registers
from docqc.checkers import structure as structure_module  # noqa: F401 - registers
from docqc.runner import check

STRUCTURE = ["structure"]
FORMULA = ["formula"]


def codes(report: object, severity: str | None = None) -> set[str]:
    """Collect finding codes from a report.

    Args:
        report: A finished :class:`~docqc.models.Report`.
        severity: When given, keep only this severity.

    Returns:
        The set of codes present.
    """
    findings = report.findings()  # type: ignore[attr-defined]
    if severity is not None:
        findings = [f for f in findings if f.severity.value == severity]
    return {f.code for f in findings}


# -- no false positives -------------------------------------------------------


def test_healthy_document_reports_nothing(fixtures: dict[str, Path]) -> None:
    """A clean document must produce zero errors and zero warnings.

    This is the most important test in the file. A checker that fires on healthy
    input trains the user to ignore it, which costs more than a missed defect.
    """
    report = check(fixtures["healthy"], names=["structure", "formula"])
    assert report.summary.errors == 0, [f.message for f in report.findings()]
    assert report.summary.warnings == 0, [f.message for f in report.findings()]
    assert report.summary.status.value == "pass"


# -- formula ------------------------------------------------------------------


def test_raw_latex_is_an_error(fixtures: dict[str, Path]) -> None:
    """Unrendered LaTeX in prose is ship-blocking."""
    report = check(fixtures["raw_latex"], names=FORMULA)
    assert "RAW_LATEX_IN_TEXT" in codes(report, "error")


def test_unbalanced_dollar_is_an_error(fixtures: dict[str, Path]) -> None:
    """An odd number of ``$`` means math was opened and never closed."""
    report = check(fixtures["unbalanced_dollar"], names=FORMULA)
    assert "UNBALANCED_MATH_DELIMITER" in codes(report, "error")


def test_duplicate_equation_number_is_an_error(fixtures: dict[str, Path]) -> None:
    """Two equations sharing a number breaks every cross-reference to it."""
    report = check(fixtures["duplicate_equation_number"], names=FORMULA)
    assert "DUPLICATE_EQUATION_NUMBER" in codes(report, "error")


def test_equation_number_gap_is_a_warning(fixtures: dict[str, Path]) -> None:
    """A skipped number usually means a deleted equation."""
    report = check(fixtures["equation_gap"], names=FORMULA)
    assert "EQUATION_NUMBER_GAP" in codes(report, "warning")


# -- structure ----------------------------------------------------------------


def test_fake_heading_is_detected(fixtures: dict[str, Path]) -> None:
    """Bold, enlarged, non-heading-styled text must be reported.

    The fixture line is unnumbered, so it lands at INFO severity — a bold line
    with no numbering might be a design choice. A *numbered* line is the
    higher-confidence case and is covered separately.
    """
    report = check(fixtures["heading_without_outline"], names=STRUCTURE)
    assert "SUSPECTED_FAKE_HEADING" in codes(report, "info")


def test_numbered_heading_without_style_is_a_warning(tmp_path: Path) -> None:
    """A numbered body line formatted as a heading is a real defect.

    Numbering signals intent: the author meant a section, so Word's TOC will be
    missing an entry. That raises it above the unnumbered case.
    """
    import docx
    from docx.shared import Pt

    doc = docx.Document()
    doc.add_paragraph("正常正文。")
    para = doc.add_paragraph()
    run_ = para.add_run("3.1 数据来源说明")
    run_.bold = True
    run_.font.size = Pt(14)
    path = tmp_path / "numbered_heading.docx"
    doc.save(path)

    report = check(path, names=STRUCTURE)
    assert "NUMBERED_HEADING_WITHOUT_STYLE" in codes(report, "warning")


def test_caption_is_not_a_fake_heading(tmp_path: Path) -> None:
    """Table and figure captions are bold and centred by design.

    Without this exemption the checker produced 94 warnings on the user's
    report, several of them ordinary captions — the false-positive failure mode
    the tool exists to avoid.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = docx.Document()
    doc.add_heading("一级标题", level=1)
    doc.add_paragraph("正文。")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ = caption.add_run("表 1-1 五维治理效能评估框架")
    run_.bold = True
    run_.font.size = Pt(10.5)
    path = tmp_path / "caption.docx"
    doc.save(path)

    report = check(path, names=STRUCTURE)
    assert "SUSPECTED_FAKE_HEADING" not in codes(report)
    assert "NUMBERED_HEADING_WITHOUT_STYLE" not in codes(report)


def test_title_block_is_not_a_fake_heading(tmp_path: Path) -> None:
    """A cover title is formatted directly and must not be reported."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = docx.Document()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ = para.add_run("低空产业治理效能评估")
    run_.bold = True
    run_.font.size = Pt(26)
    doc.add_heading("第一章 场景理解", level=1)
    doc.add_paragraph("正文。")
    path = tmp_path / "title_block.docx"
    doc.save(path)

    report = check(path, names=STRUCTURE)
    assert "SUSPECTED_FAKE_HEADING" not in codes(report)


def test_document_without_headings_is_reported(fixtures: dict[str, Path]) -> None:
    """A document with no heading at all cannot produce a TOC."""
    report = check(fixtures["no_headings"], names=STRUCTURE)
    assert "NO_HEADINGS" in codes(report, "warning")


def test_corrupt_package_reports_not_a_package(fixtures: dict[str, Path]) -> None:
    """A non-zip file gets the specific verdict, not a generic parse failure."""
    report = check(fixtures["corrupt_package"], names=STRUCTURE)
    assert "NOT_A_PACKAGE" in codes(report, "error")
    assert "UNREADABLE_DOCUMENT" not in codes(report, "error")


def test_missing_styles_part_is_reported(fixtures: dict[str, Path]) -> None:
    """A package without ``styles.xml`` cannot carry heading styles."""
    report = check(fixtures["missing_styles_part"], names=STRUCTURE)
    assert "MISSING_STYLES_PART" in codes(report, "warning")


# -- isolation ----------------------------------------------------------------


def test_one_broken_checker_does_not_stop_the_others(fixtures: dict[str, Path]) -> None:
    """A crashing checker yields ERROR while its peers still report."""
    from docqc.document import Document
    from docqc.models import CheckResult, Status
    from docqc.registry import register
    from docqc.runner import run_checks

    @register("explodes", "always raises", order=999)
    class Exploding:
        name = "explodes"
        description = "always raises"

        def run(self, document: Document) -> CheckResult:  # noqa: ARG002
            raise RuntimeError("boom")

    try:
        results = run_checks(
            Document(fixtures["healthy"]), names=["explodes", "formula"]
        )
    finally:
        from docqc.registry import _REGISTRY

        _REGISTRY.pop("explodes", None)

    by_name = {r.checker: r for r in results}
    assert by_name["explodes"].status is Status.ERROR
    assert by_name["explodes"].error is not None
    assert by_name["formula"].status is not Status.ERROR


def test_unknown_checker_name_is_rejected(fixtures: dict[str, Path]) -> None:
    """Asking for a checker that does not exist fails loudly."""
    from docqc.document import Document
    from docqc.runner import run_checks

    with pytest.raises(KeyError):
        run_checks(Document(fixtures["healthy"]), names=["does-not-exist"])
